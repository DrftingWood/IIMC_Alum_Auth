"""Idempotent loader for the IIM C alumni directory Excel file."""

from __future__ import annotations

import asyncio
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession

from src.shared.db import session_scope
from src.shared.models import Alumni

# Excel header → Alumni model attribute. Anything not listed here goes into `extra`.
COLUMN_MAP: dict[str, str] = {
    "User ID": "user_id",
    "Name": "name",
    "Email": "email",
    "Personal Email": "personal_email",
    "Status": "status",
    "Batch Year": "batch_year",
    "Program": "program",
    "Reg Number": "reg_number",
    "Mobile": "mobile",
    "Home City": "home_city",
    "Home Country": "home_country",
    "Company": "company",
    "Designation": "designation",
    "LinkedIn URL": "linkedin_url",
}


@dataclass
class ImportSummary:
    inserted: int
    updated: int


def _normalize(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        return s if s else None
    return value


def _parse_row(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any] | None:
    """Convert one Excel row into a dict suitable for Alumni insert. None if invalid."""
    record = {h: _normalize(v) for h, v in zip(headers, row, strict=False)}
    user_id = record.get("User ID")
    name = record.get("Name")
    if not user_id or not name:
        return None  # skip rows without identifier or name

    fields: dict[str, Any] = {"extra": {}}
    for header, raw in record.items():
        attr = COLUMN_MAP.get(header)
        if attr is None:
            if raw is not None:
                fields["extra"][header] = raw
            continue
        fields[attr] = raw
    return fields


async def import_directory(session: AsyncSession, path: Path) -> ImportSummary:
    """Load every row of the spreadsheet into the `alumni` table.

    Idempotent: existing rows are UPDATEd, new rows are INSERTed. Uses a
    SELECT … WHERE user_id pre-check to track accurate insert/update counts.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["All Students"] if "All Students" in wb.sheetnames else wb.active

    iterator = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(iterator)]

    inserted = 0
    updated = 0
    for raw_row in iterator:
        fields = _parse_row(headers, raw_row)
        if fields is None:
            continue
        existing = await session.get(Alumni, fields["user_id"])
        if existing is None:
            session.add(Alumni(**fields))
            inserted += 1
        else:
            for k, v in fields.items():
                if k == "user_id":
                    continue
                setattr(existing, k, v)
            updated += 1
    return ImportSummary(inserted=inserted, updated=updated)


async def _main(path: Path) -> None:
    async with session_scope() as session:
        summary = await import_directory(session, path)
        print(f"Imported: {summary.inserted} new, {summary.updated} updated.")  # noqa: T201


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python -m scripts.import_directory <path-to-xlsx>")  # noqa: T201
        sys.exit(1)
    asyncio.run(_main(Path(sys.argv[1])))
